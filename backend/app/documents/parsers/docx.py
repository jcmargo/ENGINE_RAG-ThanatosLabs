"""Parser de DOCX mediante python-docx."""

from __future__ import annotations

import zipfile
from io import BytesIO

from docx import Document

from app.documents.limits import UncompressedSizeError, assert_zip_uncompressed_within_limit
from app.documents.models import ExtractedDocument, ExtractedPage
from app.documents.parsers.base import DocumentParseError
from app.documents.versions import PARSER_VERSION

_READ_FAILED = "No se pudo leer el DOCX"


class DocxParser:
    """Lee párrafos de un DOCX y los conserva como una página."""

    def __init__(self, mime_type: str) -> None:
        self._mime_type = mime_type

    def parse(self, name: str, content: bytes) -> ExtractedDocument:
        self._assert_uncompressed_within_limit(content)
        try:
            document = Document(BytesIO(content))
            paragraphs = [
                paragraph.text.strip()
                for paragraph in document.paragraphs
                if paragraph.text.strip()
            ]
        except Exception as error:
            raise DocumentParseError(_READ_FAILED) from error
        page = ExtractedPage(number=1, text="\n\n".join(paragraphs), section=None)
        return ExtractedDocument(
            name=name,
            mime_type=self._mime_type,
            pages=(page,),
            parser_version=PARSER_VERSION,
        )

    @staticmethod
    def _assert_uncompressed_within_limit(content: bytes) -> None:
        try:
            with zipfile.ZipFile(BytesIO(content)) as archive:
                assert_zip_uncompressed_within_limit(archive)
        except UncompressedSizeError as error:
            raise DocumentParseError(error.message) from error
        except zipfile.BadZipFile as error:
            raise DocumentParseError(_READ_FAILED) from error
